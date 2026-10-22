import csv
import io
import re
from decimal import Decimal
import openpyxl

# Header Excel -> field. Beberapa nama diterima supaya format lama & baru sama-sama bisa.
# Dicocokkan lewat NAMA HEADER, bukan urutan -> kolom lain bebas diselipkan/ditambah.
ALIASES = {
    "plu": ["PLU"],
    "name": ["NAMA ITEM", "NAMA BARANG"],
    "supplier": ["NAMA SUPPLIER", "SUPPLIER"],
    "dept": ["GROUP DEPT", "DEPT"],
    "status": ["STATUS"],
    "kuu": ["KODE USAHA", "KUU", "KUU CABANG"],
    "kategori": ["KATEGORI INV", "KATEGORI"],  # opsional
    "harga": ["HARGA"],
}
# Wajib ada (kolom DAN isi tiap baris): PLU, nama, supplier, KUU, harga. Dept, status, kategori boleh kosong/tidak ada.
REQUIRED = ["plu", "name", "supplier", "kuu", "harga"]
STATUS = {"T": True, "AKTIF": True, "F": False, "TIDAK AKTIF": False}


def _price(v):
    """Harga angka atau teks. Mengerti "2,325,000.00", "1.250.000,50", "Rp 1.250.000", "1.00"."""
    if v is None or str(v).strip() == "":
        return None
    if isinstance(v, (int, float)):
        return Decimal(str(v))
    s = re.sub(r"[^\d.,]", "", str(v))
    if not re.search(r"\d", s):
        return None
    if "." in s and "," in s:
        dec = max(".", ",", key=s.rfind)  # pemisah yang muncul terakhir = desimal
    else:
        sep = "." if "." in s else "," if "," in s else None
        # satu pemisah saja: desimal bila muncul sekali & bukan 3 digit di belakangnya ("1.00", "1,5");
        # selain itu ribuan ("2,325", "1.250.000")
        dec = sep if sep and s.count(sep) == 1 and len(s.split(sep)[1]) != 3 else None
    if dec:
        whole, frac = s.rsplit(dec, 1)
        return Decimal(re.sub(r"\D", "", whole) + "." + re.sub(r"\D", "", frac))
    return Decimal(re.sub(r"\D", "", s))


def _csv_rows(data: bytes):
    try:
        text = data.decode("utf-8-sig")  # utf-8, BOM dibuang
    except UnicodeDecodeError:
        text = data.decode("cp1252")  # CSV dari Excel Windows
    first = text.splitlines()[0] if text else ""
    delim = max([",", ";", "\t", "|"], key=first.count)  # pemisah = yang paling banyak di baris header
    return csv.reader(io.StringIO(text, newline=""), delimiter=delim)


def _zero_width(fmt):
    """Lebar format nol di depan: "00000000" -> 8, "[$-409]00000000_)" -> 8, "General" -> 0."""
    sec = fmt.split(";")[0]
    sec = re.sub(r'"[^"]*"|\\.|\[[^\]]*\]|_.|\*.', "", sec)  # buang teks kutip, escape, [warna/locale], spasi
    return len(sec) if re.fullmatch(r"0+", sec) else 0


def _padded(cell):
    """Angka yang di Excel tampil dengan nol di depan (format sel 00000000) dikembalikan jadi teks bernol:
    nilai tersimpan 1234 + format 00000000 -> "00001234". Sel lain tidak diubah."""
    v = cell.value
    if isinstance(v, (int, float)) and not isinstance(v, bool) and v >= 0 and float(v).is_integer():
        s = str(int(v))
        if len(s) < (w := _zero_width(cell.number_format or "")):
            return s.zfill(w)
    return v

def _sheet_rows(file_stream):
    """Baris-baris data. .xlsx (file zip, diawali "PK") dibaca dengan openpyxl; selain itu dianggap CSV."""
    data = file_stream.read()
    if data[:2] == b"PK":
        ws = openpyxl.load_workbook(io.BytesIO(data), data_only=True).active
        return ([_padded(c) for c in row] for row in ws.iter_rows())
    return _csv_rows(data)


def load(file_stream):
    """Baca Excel master item. Mengembalikan (rows, errors)."""
    it = _sheet_rows(file_stream)
    header = [str(h).strip().upper() if h is not None else "" for h in next(it, [])]
    col = {f: header.index(n) for f, names in ALIASES.items() for n in names[::-1] if n in header}
    if missing := [f for f in REQUIRED if f not in col]:
        raise ValueError(f"Kolom tidak ditemukan di Excel: {', '.join(ALIASES[f][0] for f in missing)}")

    def raw(row, field):
        i = col.get(field)
        return row[i] if i is not None and i < len(row) else None

    def text(row, field):
        v = raw(row, field)
        if isinstance(v, float) and v.is_integer():
            v = int(v)  # PLU angka: 1001.0 -> "1001", bukan "1001.0"
        return str(v).strip() if v is not None else ""

    def opt(row, field):  # kolom tidak ada atau selnya kosong -> None -> data di DB tidak disentuh
        return text(row, field) or None

    rows, errors = [], []
    for row_idx, row in enumerate(it, start=2):
        if not any(row):
            continue
        try:
            plu, supplier, kuu, name = (text(row, f) for f in ("plu", "supplier", "kuu", "name"))
            harga = _price(raw(row, "harga"))
            required = {"PLU": plu, "Nama": name, "Supplier": supplier, "Kode Usaha": kuu, "Harga": harga}
            if empty := [k for k, v in required.items() if v in ("", None)]:
                errors.append(f"Baris {row_idx}: {', '.join(empty)} kosong atau tidak valid, dilewati.")
                continue
            status = text(row, "status").upper()
            if status and status not in STATUS:
                errors.append(f"Baris {row_idx}: Status '{text(row, 'status')}' tidak dikenali (T/F atau AKTIF/TIDAK AKTIF), dilewati.")
                continue
            rows.append({
                "plu": plu, "name": name, "supplier": supplier, "kuu": kuu, "harga": harga,
                "dept": opt(row, "dept"),
                "kategori": opt(row, "kategori"),
                "active": STATUS.get(status),  # None = status kosong/tidak ada
            })
        except Exception as e:
            errors.append(f"Baris {row_idx}: Gagal diproses ({e})")
    return rows, errors


# harga_sebelum hanya terisi saat harga benar-benar berubah.
# Baris tanpa perubahan apa pun tidak disentuh (updated_at tetap) -> "Update Hari Ini" akurat.
# dept/status/kategori yang kosong (None) tidak menimpa data yang sudah ada;
# untuk item baru: dept '' dan status Tidak Aktif.
UPSERT = """
    INSERT INTO items (plu, name, supplier, kuu, dept, active, kategori, harga_update, created_at, updated_at)
    VALUES (%(plu)s, %(name)s, %(supplier)s, %(kuu)s, COALESCE(%(dept)s, ''), COALESCE(%(active)s, FALSE),
            %(kategori)s, %(harga)s, CURRENT_TIMESTAMP, CURRENT_TIMESTAMP)
    ON CONFLICT (plu, supplier, kuu) DO UPDATE SET
        name = EXCLUDED.name,
        dept = COALESCE(%(dept)s, items.dept),
        active = COALESCE(%(active)s, items.active),
        kategori = COALESCE(%(kategori)s, items.kategori),
        harga_sebelum = CASE
            WHEN items.harga_update IS NOT NULL AND items.harga_update <> EXCLUDED.harga_update
            THEN items.harga_update ELSE items.harga_sebelum END,
        harga_update = EXCLUDED.harga_update,
        updated_at = CURRENT_TIMESTAMP
    WHERE (items.name, items.dept, items.active, items.kategori, items.harga_update)
          IS DISTINCT FROM
          (EXCLUDED.name, COALESCE(%(dept)s, items.dept), COALESCE(%(active)s, items.active),
           COALESCE(%(kategori)s, items.kategori), EXCLUDED.harga_update)
    RETURNING (xmax = 0) AS is_inserted;
"""


def upsert(conn, rows):
    inserted = updated = unchanged = 0
    with conn.cursor() as cur:
        for item in rows:
            cur.execute(UPSERT, item)
            res = cur.fetchone()
            if res is None:
                unchanged += 1
            elif res[0]:
                inserted += 1
            else:
                updated += 1
    return {"inserted": inserted, "updated": updated, "unchanged": unchanged}


if __name__ == "__main__":  # cek cepat parser: python importer.py
    def mk(rows):
        wb = openpyxl.Workbook(); ws = wb.active
        for r in rows:
            ws.append(r)
        b = io.BytesIO(); wb.save(b); b.seek(0)
        return load(b)

    r, e = mk([["PLU", "HARGA", "NAMA BARANG", "SUPPLIER", "GROUP DEPT", "STATUS", "KUU"],
               ["01001", "Rp 1.250.000,50", "A", "S", "OSA", "AKTIF", "SCZ1"],
               ["0002", 5000, "B", "S", "OSA", "TIDAK AKTIF", "SCZ1"],
               ["0003", None, "C", "S", "OSA", "AKTIF", "SCZ1"],        # harga kosong -> error
               ["0004", 1, "", "S", "OSA", "AKTIF", "SCZ1"]])           # nama kosong -> error
    assert [x["harga"] for x in r] == [Decimal("1250000.50"), Decimal("5000")], r
    assert r[0]["plu"] == "01001" and r[1]["active"] is False and len(e) == 2, e
    # hanya kolom wajib: dept/status/kategori tidak ada -> None (DB tidak disentuh)
    r, e = mk([["PLU", "NAMA ITEM", "NAMA SUPPLIER", "KODE USAHA", "HARGA"], ["1", "A", "S", "K", 10]])
    assert not e and (r[0]["dept"], r[0]["active"], r[0]["kategori"]) == (None, None, None), r
    # kolom wajib hilang -> ditolak
    try:
        mk([["PLU", "NAMA ITEM", "NAMA SUPPLIER", "KODE USAHA"], ["1", "A", "S", "K"]])
        raise SystemExit("harusnya ValueError")
    except ValueError:
        pass
    for txt, want in [("2,325,000.00", "2325000.00"), ("1.00", "1.00"), ("500,000.00", "500000.00"),
                      ("Rp 1.250.000", "1250000"), ("1.250.000,50", "1250000.50"), ("2,325", "2325"),
                      ("1,5", "1.5"), ("-", None), ("", None)]:
        got = _price(txt)
        assert got == (Decimal(want) if want else None), (txt, got)
    print("ok")